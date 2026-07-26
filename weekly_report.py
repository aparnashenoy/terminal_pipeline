"""
weekly_report.py

Automated weekly reporting pipeline for port terminal operations.

Usage:
    python3 weekly_report.py

What it does (classic ETL):
    1. EXTRACT  - reads raw_terminal_data.csv (as if pulled from a TOS)
    2. TRANSFORM - filters to the last 7 days, computes KPIs
    3. LOAD     - writes a formatted Excel report to /output

This is the script that, in production, would later be wrapped in an
Airflow DAG and scheduled to run automatically every Monday morning.
"""

import logging
import sys
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from notifier import send_report

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
INPUT_FILE = Path("raw_terminal_data.csv")
OUTPUT_DIR = Path("output")
REPORT_DAYS = 7  # "weekly" window

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    handlers=[
        logging.FileHandler("pipeline.log"),
        logging.StreamHandler(sys.stdout),
    ],
)
log = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# 1. EXTRACT
# ---------------------------------------------------------------------------
def extract(input_file: Path) -> pd.DataFrame:
    log.info(f"Extracting data from {input_file}")
    if not input_file.exists():
        log.error(f"Input file not found: {input_file}")
        raise FileNotFoundError(f"Cannot find {input_file}. Run generate_sample_data.py first.")

    df = pd.read_csv(input_file, parse_dates=["date"])
    log.info(f"Loaded {len(df)} rows")
    return df


# ---------------------------------------------------------------------------
# 2. TRANSFORM
# ---------------------------------------------------------------------------
def transform(df: pd.DataFrame, report_days: int = 7):
    latest_date = df["date"].max()
    window_start = latest_date - timedelta(days=report_days - 1)

    week_df = df[(df["date"] >= window_start) & (df["date"] <= latest_date)].copy()
    log.info(
        f"Filtered to reporting window {window_start.date()} -> {latest_date.date()} "
        f"({len(week_df)} vessel calls)"
    )

    if week_df.empty:
        log.warning("No data found in the reporting window.")

    # --- Overall summary KPIs ---
    summary = {
        "Total Vessel Calls": int(week_df["vessel_call_id"].nunique()),
        "Total Containers Handled": int(week_df["containers_handled"].sum()),
        "Total TEU Moves": int(week_df["teu_moves"].sum()),
        "Avg Crane Productivity (moves/hr)": round(week_df["crane_moves_per_hour"].mean(), 1),
        "Avg Berth Hours per Call": round(week_df["berth_hours"].mean(), 1),
        "Avg Dwell Time (hrs)": round(week_df["avg_dwell_time_hours"].mean(), 1),
    }

    # --- By berth ---
    by_berth = (
        week_df.groupby("berth")
        .agg(
            vessel_calls=("vessel_call_id", "nunique"),
            teu_moves=("teu_moves", "sum"),
            avg_berth_hours=("berth_hours", "mean"),
            avg_dwell_hours=("avg_dwell_time_hours", "mean"),
        )
        .round(1)
        .sort_values("teu_moves", ascending=False)
        .reset_index()
    )

    # --- By shipping line ---
    by_line = (
        week_df.groupby("shipping_line")
        .agg(
            vessel_calls=("vessel_call_id", "nunique"),
            teu_moves=("teu_moves", "sum"),
            avg_dwell_hours=("avg_dwell_time_hours", "mean"),
        )
        .round(1)
        .sort_values("teu_moves", ascending=False)
        .reset_index()
    )

    # --- By operation type ---
    by_operation = (
        week_df.groupby("operation_type")
        .agg(
            vessel_calls=("vessel_call_id", "nunique"),
            teu_moves=("teu_moves", "sum"),
        )
        .round(1)
        .sort_values("teu_moves", ascending=False)
        .reset_index()
    )

    log.info("Transform complete: summary + 3 breakdown tables built")
    return summary, by_berth, by_line, by_operation, week_df, window_start, latest_date


# ---------------------------------------------------------------------------
# 3. LOAD  (write formatted Excel report)
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Arial", bold=True, size=14, color="1F4E78")
LABEL_FONT = Font(name="Arial", bold=True, size=11)
BODY_FONT = Font(name="Arial", size=10)


def _style_header_row(ws, row_idx, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autofit(ws, num_cols, min_width=12):
    for col in range(1, num_cols + 1):
        letter = get_column_letter(col)
        max_len = max(
            (len(str(cell.value)) for cell in ws[letter] if cell.value is not None),
            default=min_width,
        )
        ws.column_dimensions[letter].width = max(min_width, max_len + 2)


def write_dataframe(ws, df: pd.DataFrame, start_row: int, title: str = None):
    row = start_row
    if title:
        ws.cell(row=row, column=1, value=title).font = LABEL_FONT
        row += 1

    for col_idx, col_name in enumerate(df.columns, start=1):
        ws.cell(row=row, column=col_idx, value=col_name)
    _style_header_row(ws, row, len(df.columns))
    row += 1

    for _, record in df.iterrows():
        for col_idx, value in enumerate(record, start=1):
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.font = BODY_FONT
        row += 1

    return row + 1  # next free row


def load(summary, by_berth, by_line, by_operation, window_start, latest_date, output_dir: Path):
    output_dir.mkdir(exist_ok=True)
    report_date_str = latest_date.strftime("%Y-%m-%d")
    output_path = output_dir / f"weekly_terminal_report_{report_date_str}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Weekly Report"

    ws.cell(row=1, column=1, value="Weekly Terminal Operations Report").font = TITLE_FONT
    ws.cell(
        row=2, column=1,
        value=f"Reporting window: {window_start.date()} to {latest_date.date()}",
    ).font = Font(name="Arial", italic=True, size=10)

    row = 4
    ws.cell(row=row, column=1, value="Key Metrics").font = LABEL_FONT
    row += 1
    for label, value in summary.items():
        ws.cell(row=row, column=1, value=label).font = BODY_FONT
        ws.cell(row=row, column=2, value=value).font = BODY_FONT
        row += 1
    row += 1

    row = write_dataframe(ws, by_berth, row, title="Performance by Berth")
    row = write_dataframe(ws, by_line, row, title="Performance by Shipping Line")
    row = write_dataframe(ws, by_operation, row, title="Performance by Operation Type")

    _autofit(ws, max(6, len(by_line.columns)))

    wb.save(output_path)
    log.info(f"Report written to {output_path}")
    return output_path


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    try:
        df = extract(INPUT_FILE)
        summary, by_berth, by_line, by_operation, week_df, window_start, latest_date = transform(
            df, REPORT_DAYS
        )
        output_path = load(summary, by_berth, by_line, by_operation, window_start, latest_date, OUTPUT_DIR)

        # dry_run=True here since no real SMTP credentials exist in this environment.
        # On your own machine, set the SMTP_* env vars and switch to dry_run=False.
        email_sent = send_report(output_path, summary, dry_run=True)
        if not email_sent:
            log.warning("Report generated but email delivery was skipped/failed. See log above.")

        log.info("Pipeline completed successfully.")
        print(f"\nReport generated: {output_path}")
    except Exception:
        log.exception("Pipeline failed")
        sys.exit(1)


if __name__ == "__main__":
    main()
